"""导出服务：将 InterviewResult 导出为 Word 或 PDF."""

import os
from datetime import datetime
from io import BytesIO
from typing import Optional

from ..models.schemas import InterviewQuestions, InterviewResult, Question


def _find_cjk_font() -> Optional[str]:
    """查找系统中可用的中文字体，用于 PDF 导出."""
    candidates = [
        # macOS
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/Hiragino Sans GB.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        # Linux
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        # Windows
        "C:/Windows/Fonts/msyh.ttc",
        "C:/Windows/Fonts/simsun.ttc",
        "C:/Windows/Fonts/simhei.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def _format_questions(questions: InterviewQuestions) -> list[tuple[str, list[Question]]]:
    """将五方向题目整理为顺序列表."""
    return [
        ("方向一：简历深挖", questions.resume_deep_dive),
        ("方向二：业务场景题", questions.business_scenario),
        ("方向三：技术/专业题", questions.technical_professional),
        ("方向四：行为/文化题", questions.behavioral_cultural),
        ("方向五：压力/应变题", questions.pressure_adaptability),
    ]


def generate_word(result: InterviewResult) -> bytes:
    """生成 Word 文档（.docx）字节流."""
    try:
        from docx import Document
    except ImportError as exc:
        raise ImportError("导出 Word 需要安装 python-docx：pip install python-docx") from exc

    doc = Document()

    # 标题
    title = doc.add_heading("HireMind 面试分析报告", level=0)
    title.alignment = 1  # 居中
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"目标公司：{result.company_name}")
    doc.add_paragraph(f"目标岗位：{result.job_title}")

    # 匹配度分析
    doc.add_heading("一、匹配度分析", level=1)
    doc.add_heading("强匹配", level=2)
    for item in result.match_analysis.strong_points:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("弱项 / 需补强", level=2)
    for item in result.match_analysis.weak_points:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("JD 要求但简历未提", level=2)
    for item in result.match_analysis.not_mentioned:
        doc.add_paragraph(item, style="List Bullet")

    # 业务场景推演
    doc.add_heading("二、业务场景推演", level=1)
    for idx, s in enumerate(result.business_scenarios, start=1):
        doc.add_heading(f"场景 {idx}：{s.scenario}", level=2)
        doc.add_paragraph(f"可能问题：{s.inferred_problem}")
        doc.add_paragraph(f"面试切入：{s.interview_angle}")

    # 风险点识别
    doc.add_heading("三、风险点识别", level=1)
    for risk in result.risks:
        doc.add_heading(f"{risk.risk_type}（严重度：{risk.severity}）", level=2)
        doc.add_paragraph(f"描述：{risk.description}")
        doc.add_paragraph(f"建议追问：{risk.suggested_question}")

    # 五方向面试题
    doc.add_heading("四、五方向面试题", level=1)
    for direction_title, qs in _format_questions(result.questions):
        doc.add_heading(direction_title, level=2)
        for idx, q in enumerate(qs, start=1):
            doc.add_heading(f"Q{idx}：{q.title}（{q.difficulty}）", level=3)
            doc.add_paragraph(f"题目：{q.question}")
            doc.add_paragraph(f"考察点：{q.focus}")
            doc.add_paragraph(f"参考提示：{q.reference_hint}")
            if q.follow_ups:
                doc.add_paragraph("追问方向：")
                for fu in q.follow_ups:
                    doc.add_paragraph(fu, style="List Bullet")

    # 备考建议
    if result.preparation_tips:
        doc.add_heading("五、备考建议", level=1)
        for tip in result.preparation_tips:
            doc.add_paragraph(tip, style="List Bullet")

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_pdf(result: InterviewResult) -> bytes:
    """生成 PDF 文档字节流."""
    try:
        from fpdf import FPDF, XPos, YPos
    except ImportError as exc:
        raise ImportError("导出 PDF 需要安装 fpdf2：pip install fpdf2") from exc

    font_path = _find_cjk_font()
    if not font_path:
        raise RuntimeError(
            "未找到系统中文字体，无法生成中文 PDF。"
            "请安装中文字体，或改用 Word 导出。"
        )

    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("CJK", "", font_path)
    pdf.set_font("CJK", size=12)

    def _section(title: str):
        pdf.set_font_size(16)
        pdf.cell(0, 10, text=title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font_size(12)
        pdf.ln(2)

    def _subsection(title: str):
        pdf.set_font_size(13)
        pdf.cell(0, 8, text=title, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.set_font_size(12)

    def _text(text: str):
        pdf.multi_cell(0, 6, text=text)
        pdf.ln(1)

    # 标题
    pdf.set_font_size(20)
    pdf.cell(0, 12, text="HireMind 面试分析报告", align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font_size(12)
    _text(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}")
    _text(f"目标公司：{result.company_name}")
    _text(f"目标岗位：{result.job_title}")

    # 匹配度分析
    _section("一、匹配度分析")
    _subsection("强匹配")
    for item in result.match_analysis.strong_points:
        _text(f"• {item}")
    _subsection("弱项 / 需补强")
    for item in result.match_analysis.weak_points:
        _text(f"• {item}")
    _subsection("JD 要求但简历未提")
    for item in result.match_analysis.not_mentioned:
        _text(f"• {item}")

    # 业务场景推演
    _section("二、业务场景推演")
    for idx, s in enumerate(result.business_scenarios, start=1):
        _subsection(f"场景 {idx}：{s.scenario}")
        _text(f"可能问题：{s.inferred_problem}")
        _text(f"面试切入：{s.interview_angle}")

    # 风险点识别
    _section("三、风险点识别")
    for risk in result.risks:
        _subsection(f"{risk.risk_type}（严重度：{risk.severity}）")
        _text(f"描述：{risk.description}")
        _text(f"建议追问：{risk.suggested_question}")

    # 五方向面试题
    _section("四、五方向面试题")
    for direction_title, qs in _format_questions(result.questions):
        _subsection(direction_title)
        for idx, q in enumerate(qs, start=1):
            _text(f"Q{idx}：{q.title}（{q.difficulty}）")
            _text(f"  题目：{q.question}")
            _text(f"  考察点：{q.focus}")
            _text(f"  参考提示：{q.reference_hint}")
            if q.follow_ups:
                for fu in q.follow_ups:
                    _text(f"  • 追问：{fu}")

    # 备考建议
    if result.preparation_tips:
        _section("五、备考建议")
        for tip in result.preparation_tips:
            _text(f"• {tip}")

    return bytes(pdf.output())
